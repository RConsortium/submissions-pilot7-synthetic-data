"""Evidence-backed metadata consistency checks and browsable source material.

Narrative interpretation is explicitly incomplete. Automated findings are
observations for human review, not a regulatory or clinical judgement.
"""

import html
import json
from pathlib import Path
import re
import shutil
from urllib.parse import quote
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader
import pyarrow.parquet as pq


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=True, separators=(",", ":")) + "\n")


def local_name(element):
    return element.tag.rsplit("}", 1)[-1]


def description(element):
    text = " ".join(
        " ".join("".join(x.itertext()).split())
        for x in element
        if local_name(x) == "Description"
    ).strip()
    # Define-XML 1.0 stores the label in def:Label, not Description.
    return text or next(
        (v for k, v in element.attrib.items() if k.rsplit("}", 1)[-1] == "Label"), ""
    )


def define_metadata(path):
    tree = ET.parse(path)
    items = {x.attrib["OID"]: x for x in tree.iter() if local_name(x) == "ItemDef"}
    methods = {
        x.attrib["OID"]: description(x)
        for x in tree.iter()
        if local_name(x) == "MethodDef"
    }
    groups = {}
    for group in tree.iter():
        if local_name(group) != "ItemGroupDef":
            continue
        fields = {}
        for reference in group:
            if local_name(reference) != "ItemRef":
                continue
            item = items.get(reference.attrib.get("ItemOID"))
            if item is None:
                continue
            fields[item.attrib["Name"]] = {
                "label": description(item),
                "type": item.attrib.get("DataType"),
                "oid": item.attrib["OID"],
                "method": methods.get(reference.attrib.get("MethodOID"), ""),
            }
        groups[group.attrib.get("Name", "")] = fields
    return groups


def document_html(title, body):
    return (
        '<!doctype html><html lang="en"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>{html.escape(title)}</title><style>"
        "body{font:16px/1.7 system-ui;margin:30px auto;padding:0 24px;max-width:1000px;color:#24344a}"
        "section{border-bottom:1px solid #ddd;padding:24px 0;scroll-margin-top:20px}"
        "pre{white-space:pre-wrap;font:inherit}td,th{padding:8px;border:1px solid #ddd;vertical-align:top}"
        "table{border-collapse:collapse}a{color:#2d5f8a}:target{background:#fff5d6}"
        "</style><body>" + f"<h1>{html.escape(title)}</h1>" + body + "</body></html>"
    )


def materials(root, study, target, metadata):
    folder = target / "documents"
    folder.mkdir(exist_ok=True)
    manifest_path = root / study / "spec/source/manifest.json"
    manifest = (
        json.loads(manifest_path.read_text())
        if manifest_path.exists()
        else {"files": []}
    )
    documents = []
    for entry in manifest["files"]:
        path = root / entry["path"]
        if not path.exists():
            continue
        original = folder / path.name
        shutil.copyfile(path, original)
        record = {
            "name": path.name,
            "path": entry["path"],
            "role": entry.get("role"),
            "upstream": entry.get("upstream", {}),
            "original": f"documents/{path.name}",
            "sections": [],
        }
        if path.suffix.lower() == ".pdf":
            for number, page in enumerate(PdfReader(path).pages, 1):
                record["sections"].append(
                    {
                        "id": f"page-{number}",
                        "title": f"Page {number}",
                        "text": page.extract_text() or "",
                    }
                )
        elif path.suffix.lower() == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=True)
            for sheet in book:
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(value) if value is not None else "" for value in row]
                    if any(values):
                        record["sections"].append(
                            {
                                "id": f'{quote(sheet.title, safe="")}-row-{number}',
                                "title": f"{sheet.title} / row {number}",
                                "text": " | ".join(values),
                            }
                        )
            book.close()
        elif path.name == "define-adam.xml":
            for domain, columns in metadata.items():
                for name, item in columns.items():
                    record["sections"].append(
                        {
                            "id": f"{domain}.{name}",
                            "title": f'{domain}.{name} / {item["oid"]}',
                            "text": f'Label: {item["label"]}\nData type: {item["type"]}\nMethod: {item["method"]}',
                        }
                    )
        if record["sections"]:
            body = "<p>Extracted reading copy. Verify against the linked original submission document.</p>"
            body += f'<p><a href="{html.escape(path.name)}">Open original {html.escape(path.name)}</a></p>'
            for section in record["sections"]:
                body += (
                    f'<section id="{html.escape(section["id"])}"><h2>{html.escape(section["title"])}</h2>'
                    f'<pre>{html.escape(section["text"])}</pre></section>'
                )
            reading = f"{path.stem}.html"
            (folder / reading).write_text(
                document_html(path.name, body), encoding="utf-8"
            )
            record["reading"] = f"documents/{reading}"
        documents.append(record)
    save_json(target / "documents.json", documents)
    return manifest, documents


def build_report(root, study, specifications, target, revision):
    define_path = root / study / "spec/define/define-adam.xml"
    metadata = define_metadata(define_path) if define_path.exists() else {}
    manifest, documents = materials(root, study, target, metadata)
    findings, checks = [], {"labels": 0, "output_columns": 0, "define_variables": 0}
    repo = "https://github.com/RConsortium/submissions-pilot7-synthetic-data"

    def add(
        domain,
        variable,
        kind,
        title,
        observation,
        source,
        spec_path,
        line,
        recommendation,
    ):
        findings.append(
            {
                "id": f'{study}:{domain}:{variable or "study"}:{kind}',
                "dataset": domain,
                "variable": variable,
                "kind": kind,
                "status": "Needs review",
                "title": title,
                "observation": observation,
                "source": source,
                "spec": {
                    "path": str(spec_path.relative_to(root)),
                    "line": line,
                    "url": f"{repo}/blob/{revision}/{spec_path.relative_to(root)}#L{line}",
                },
                "recommendation": recommendation,
            }
        )

    for spec_path, spec, lineage in specifications:
        domain = spec["domain"]
        output_path = (spec_path.parent / spec["output"]["path"]).resolve()
        output_columns = set(pq.read_schema(output_path).names)
        declared = {item["name"]: item for item in spec["columns"]}
        published = metadata.get(domain, {})
        if not published:
            # Missing name is a question; aliases must be verified by a reviewer.
            add(
                domain,
                "",
                "unclear",
                "Confirm dataset naming against Define-XML",
                f"The specification declares {domain}; no ItemGroupDef with that exact name was found. "
                f'Published names include {", ".join(metadata)}.',
                {
                    "path": str(define_path.relative_to(root)),
                    "section": "Dataset inventory",
                    "excerpt": ", ".join(metadata),
                    "url": "documents/define-adam.xml",
                },
                spec_path,
                1,
                "Verify the intended dataset alias in the published package and existing comparator.",
            )
        for name in spec["output"]["columns"]:
            checks["output_columns"] += 1
            line = lineage["spans"].get(name, {}).get("start", 1)
            if name not in output_columns:
                add(
                    domain,
                    name,
                    "inconsistent",
                    "Declared output column is absent",
                    f"{name} is declared in output.columns but is absent from the derived dataset.",
                    {
                        "path": str(output_path.relative_to(root)),
                        "section": "Dataset schema",
                        "excerpt": ", ".join(sorted(output_columns)),
                        "url": "",
                    },
                    spec_path,
                    line,
                    "Review the output selection and build evidence.",
                )
            if name not in published:
                continue
            checks["define_variables"] += 1
            item, column = published[name], declared.get(name, {})
            checks["labels"] += 1
            if column.get("label", "").strip() != item["label"].strip():
                add(
                    domain,
                    name,
                    "inconsistent",
                    "Variable label differs from published Define-XML",
                    f'YAML label: {column.get("label", "(missing)")}. '
                    f'Published label: {item["label"]}.',
                    {
                        "path": str(define_path.relative_to(root)),
                        "section": f'{domain}.{name} / {item["oid"]}',
                        "excerpt": f'Label: {item["label"]}\nMethod: {item["method"]}',
                        "url": f"documents/define-adam.html#{domain}.{name}",
                    },
                    spec_path,
                    line,
                    "Check the original ItemDef and determine whether the difference is intended.",
                )
    if manifest.get("reference_data_alignment") == "not_revalidated":
        spec_path, spec, _ = specifications[0]
        notes = "\n".join(manifest.get("alignment_notes", []))
        add(
            spec["domain"],
            "",
            "unclear",
            "Confirm submission-document version alignment",
            "The retained source manifest records reference_data_alignment as not_revalidated.",
            {
                "path": f"{study}/spec/source/manifest.json",
                "section": "reference_data_alignment / alignment_notes",
                "excerpt": notes or "reference_data_alignment: not_revalidated",
                "url": f"{repo}/blob/{revision}/{study}/spec/source/manifest.json",
            },
            spec_path,
            1,
            "Verify which package version applies before interpreting narrative rules.",
        )
    # A real documented rule that can be opened alongside its implementation.
    # It is a review task, never reported as an observed inconsistency.
    guide = next(
        (d for d in documents if d["name"] in {"dataguide.pdf", "adrg.pdf"}), None
    )
    if guide:
        section = next(
            (
                s
                for s in guide["sections"]
                if "imput" in s["text"].lower()
                and any(word in s["text"].upper() for word in ["ADAE", "ASTDT"])
            ),
            None,
        )
        adae = next(
            (item for item in specifications if item[1]["domain"] == "ADAE"), None
        )
        if section and adae:
            spec_path, _, lineage = adae
            text = section["text"]
            match = text.lower().find("imput")
            start = text.rfind("\n", 0, max(0, match - 180)) + 1
            end = text.find("\n", min(len(text), match + 650))
            excerpt = text[start:end if end >= 0 else len(text)]
            add(
                "ADAE",
                "ASTDT",
                "not_reviewed",
                "Review date imputation against the analysis guide",
                "The guide describes date imputation. A clinical interpretation review of the implemented rule is pending.",
                {
                    "path": guide["path"],
                    "section": section["title"],
                    "excerpt": excerpt,
                    "url": f'{guide["reading"]}#{section["id"]}',
                    "original": guide["original"]
                    + "#page="
                    + section["id"].removeprefix("page-"),
                },
                spec_path,
                lineage["spans"].get("ASTDT", {}).get("start", 1),
                "Compare the quoted rule with the ASTDT specification, partial source dates and derived dates. Record your interpretation in a comment.",
            )
    return {
        "study": study,
        "revision": revision,
        "findings": findings,
        "checks": checks,
        "documents": [
            {k: v for k, v in d.items() if k != "sections"} for d in documents
        ],
        "coverage": {
            "metadata": "Automated output-presence and exact Define-XML label checks completed.",
            "narrative": "Not comprehensively reviewed; guide review tasks are explicitly marked Not reviewed.",
            "clinical_judgement": "Requires human review.",
        },
        "note": "Automated observations and pending review tasks are not confirmed clinical defects.",
    }
