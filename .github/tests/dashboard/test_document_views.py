"""Original document views retain structure, escape text and confine XSLT reads."""

import json
from pathlib import Path
import sys
import tempfile
import unittest

from lxml import etree
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
from study_dashboard.document_views import render_xml, workbook_body
from study_dashboard.documents import materials, program_view


class DocumentViewTests(unittest.TestCase):
    def test_protocol_conversion_keeps_original_bytes_and_escapes_reading_view(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / 'study/spec/source/protocol-sap.mmd'
            source.parent.mkdir(parents=True)
            text = 'Age \u2264 65; <script>retained source</script>\n'
            source.write_text(text, encoding='utf-8')
            (source.parent / 'manifest.json').write_text(
                '{"files":[{"path":"study/spec/source/protocol-sap.mmd"}]}')
            target = root / 'bundle'
            target.mkdir()
            _, documents = materials(root, 'study', target, {})
            document = documents[0]
            self.assertEqual((target / document['original']).read_bytes(), source.read_bytes())
            view = (target / document['view']).read_text(encoding='ascii')
            tree = etree.HTML(view)
            self.assertEqual(tree.xpath('string(//pre)'), text)
            self.assertEqual(tree.xpath('//script'), [])

    def test_program_view_numbers_escaped_lines_and_names_its_provenance(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / 'study/spec/source/programs/adlbc.r'
            source.parent.mkdir(parents=True)
            text = 'ONE <- abs(x)\nif (a < b) "</pre>"\n'
            source.write_text(text, encoding='ascii')
            upstream = {'repository': 'o/r', 'commit': '0123456789abcdef', 'path': 'programs/adlbc.r',
                        'url': 'https://github.com/o/r/blob/0123456789abcdef/programs/adlbc.r'}
            (root / 'study/spec/source/manifest.json').write_text(json.dumps({'files': [
                {'path': 'study/spec/source/programs/adlbc.r', 'role': 'reference_reimplementation', 'upstream': upstream}]}))
            target = root / 'bundle'
            target.mkdir()
            _, documents = materials(root, 'study', target, {})
            document = documents[0]
            self.assertEqual(
                {k: document[k] for k in ('view', 'reading', 'view_label', 'sections')},
                {'view': 'documents/study/spec/source/programs/adlbc.r.html',
                 'reading': 'documents/study/spec/source/programs/adlbc.r.html',
                 'view_label': 'Program lines', 'sections': []})
            self.assertEqual((target / document['original']).read_bytes(), source.read_bytes())
            tree = etree.HTML((target / document['view']).read_text(encoding='utf-8'))
            self.assertEqual(
                [(line.get('id'), line.xpath('string(.)')) for line in tree.xpath('//pre/span[@class="line"]')],
                [('L1', '1 ONE <- abs(x)\n'), ('L2', '2 if (a < b) "</pre>"\n')])
            self.assertEqual(tree.xpath('string(//a[@id="upstream"]/@href)'), upstream['url'])
            self.assertIn('not what the original authors intended', tree.xpath('string(//p[1])'))
            self.assertEqual(tree.xpath('//pre//span[@class="kr"]/text()'), ['if'])
            self.assertEqual(tree.xpath('//pre//span[@class="o"]/text()'), ['<-', '<'])
            self.assertEqual(tree.xpath('//pre//span[@class="s"]/text()'), ['"</pre>"'])
            self.assertEqual(tree.xpath('//pre//script | //script[@src] | //link[@href]'), [])

    def test_r_highlighting_preserves_blank_lines_tabs_and_multiline_strings(self):
        cases = [
            ('empty.r', ''),
            ('blank.R', '\n\n'),
            ('adsl.r', '\n# <script>source</script>\nx <- "first\n\tsecond & last"\n\n'),
            ('ADSL.R', '\tx <- 42 # no final newline'),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            for name, text in cases:
                with self.subTest(name=name):
                    path = Path(tmp) / name
                    path.write_text(text, encoding='ascii')
                    tree = etree.HTML(program_view(path, {}))
                    lines = text.splitlines()
                    width = len(str(len(lines)))
                    self.assertEqual(
                        [(line.get('id'), line.xpath('string(.)'))
                         for line in tree.xpath('//pre/span[@class="line"]')],
                        [(f'L{n}', f'{n:>{width}} {line}\n')
                         for n, line in enumerate(lines, 1)])
                    self.assertEqual(tree.xpath('//pre//script'), [])
                    if name == 'adsl.r':
                        self.assertEqual(tree.xpath('//pre//span[@class="c1"]/text()'),
                                         ['# <script>source</script>'])
                        self.assertEqual(tree.xpath('//pre//span[@class="s"]/text()'),
                                         ['"first', '\tsecond & last"'])
                    if name == 'ADSL.R':
                        self.assertEqual(tree.xpath('//pre//span[@class="m"]/text()'), ['42'])

    def test_other_programs_keep_escaped_plain_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'adsl.sas'
            path.write_text('if a < b then label = "</pre>";\n', encoding='ascii')
            tree = etree.HTML(program_view(path, {}))
            self.assertEqual(tree.xpath('string(//pre)'), '1 if a < b then label = "</pre>";\n')
            self.assertEqual(tree.xpath('//pre//span/@class'), ['line', 'line-number'])
            self.assertEqual(tree.xpath('//pre//script'), [])

    def test_workbook_keeps_coordinates_merges_formulas_and_escaped_cell_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "spec.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Rules & inputs"
            sheet["A1"] = "<script>source text</script>"
            sheet.merge_cells("A1:B1")
            sheet["B3"] = "=1+2"
            book.create_sheet("Output")["A1"] = "ADAE"
            book.save(path)
            body, sections = workbook_body(path)
            tree = etree.HTML(body)
            self.assertEqual(
                tree.xpath('string(//tr[@id="sheet-1-row-1"]/td)'),
                "<script>source text</script>",
            )
            self.assertEqual(tree.xpath('//tr[@id="sheet-1-row-1"]/td/@colspan'), ["2"])
            self.assertEqual(
                tree.xpath('//tr[@id="sheet-1-row-3"]/td/text()'), ["=1+2"]
            )
            self.assertEqual(tree.xpath("//script"), [])
            self.assertEqual(
                [s["id"] for s in sections],
                ["sheet-1-row-1", "sheet-1-row-3", "sheet-2-row-1"],
            )
            self.assertEqual(
                tree.xpath("//caption/text()"), ["Rules & inputs", "Output"]
            )
            self.assertEqual(tree.xpath('//tbody/tr/@id'), [s['id'] for s in sections])
            self.assertEqual(tree.xpath('//button[@role="tab"]/text()'), ['Rules & inputs', 'Output'])
            self.assertEqual(tree.xpath('//button[@aria-selected="true"]/@aria-controls'), ['sheet-1'])
            self.assertEqual(tree.xpath('//section[@role="tabpanel" and not(@hidden)]/@id'), ['sheet-1'])
            self.assertEqual(tree.xpath('//section[@role="tabpanel"]//table/@class'), ['review-table', 'review-table'])

    def test_workbook_omits_blank_rows_and_keeps_merges_aligned(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / 'study/spec/source/spec.xlsx'
            source.parent.mkdir(parents=True)
            book = Workbook()
            sheet = book.active
            sheet.append([0, False])
            sheet['A2'] = ' \t '
            sheet['A3'] = '=1+2'
            sheet.merge_cells('A3:A5')
            sheet['B5'] = 'last merged row'
            sheet.merge_cells('A7:A9')
            sheet['B7'] = 'empty merged cell'
            sheet['B9'] = 'last empty merged row'
            sheet['D10'].number_format = '0.00'
            book.save(source)
            (source.parent / 'manifest.json').write_text(
                '{"files":[{"path":"study/spec/source/spec.xlsx"}]}')
            target = root / 'bundle'
            target.mkdir()
            _, documents = materials(root, 'study', target, {})
            document = documents[0]
            tree = etree.HTML((target / document['view']).read_bytes())
            self.assertEqual(tree.xpath('//tbody/tr/@id'), [
                f'sheet-1-row-{number}' for number in (1, 3, 5, 7, 9)
            ])
            self.assertEqual(tree.xpath('//tr[@id="sheet-1-row-1"]/td/text()'), ['0', 'False'])
            self.assertEqual(tree.xpath('//td[@rowspan="2"]/text()'), ['=1+2'])
            self.assertEqual(tree.xpath('//tr[@id="sheet-1-row-7"]/td/@rowspan'), ['2', '1', '1', '1'])
            self.assertEqual(tree.xpath('//tr[@id="sheet-1-row-9"]/td/text()'), ['last empty merged row'])
            self.assertEqual((target / document['original']).read_bytes(), source.read_bytes())
            self.assertEqual(tree.xpath('//button[@id="compare-worksheets"]'), [])
            self.assertEqual(len(tree.xpath('//script')), 1)

    def test_empty_worksheet_remains_selectable_and_sheet_titles_are_escaped(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'empty.xlsx'
            book = Workbook()
            book.active.title = 'Empty <&>'
            book.save(path)
            body, sections = workbook_body(path)
            tree = etree.HTML(body)
            self.assertEqual(sections, [])
            self.assertEqual(tree.xpath('//button[@role="tab"]/text()'), ['Empty <&>'])
            self.assertEqual(tree.xpath('//section[@id="sheet-1"]//p/text()'), ['This worksheet is empty.'])

    def test_xml_uses_retained_publisher_html_and_denies_external_document_reads(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "define.xml"
            source.write_text('<ODM><ItemDef Name="ASTDT"/></ODM>')
            stylesheet = root / "original.xsl"
            template = '<xsl:stylesheet version="1.0" xmlns:xsl="http://www.w3.org/1999/XSL/Transform"><xsl:output method="html"/><xsl:template match="/"><html><body><table id="published"><tr><td><xsl:value-of select="ODM/ItemDef/@Name"/></td></tr></table>EXTERNAL</body></html></xsl:template></xsl:stylesheet>'
            stylesheet.write_text(template.replace("EXTERNAL", ""))
            view = render_xml(root, source, source, {"stylesheet": "original.xsl"})
            self.assertEqual(
                etree.HTML(view.read_bytes()).xpath(
                    '//table[@id="published"]//td/text()'
                ),
                ["ASTDT"],
            )
            self.assertEqual(source.read_text(), '<ODM><ItemDef Name="ASTDT"/></ODM>')
            # A cited OID, exact or prefixed by the stylesheet, highlights its row inside the body.
            html = view.read_text()
            self.assertIn("document.getElementById(want)||all.find(function(e){return e.id.endsWith('.IT.'+want)})", html)
            self.assertIn("hit=hit.closest('tr')||hit;hit.classList.add('yamaa-cited')", html)
            self.assertIn("tr.yamaa-cited>*{background:", html)
            self.assertLess(html.index("<script>"), html.lower().index("</body>"))
            stylesheet.write_text(
                template.replace(
                    "EXTERNAL", "<xsl:value-of select=\"document('define.xml')\"/>"
                )
            )
            with self.assertRaises(etree.XSLTApplyError):
                render_xml(root, source, source, {"stylesheet": "original.xsl"})
            with self.assertRaises(ValueError):
                render_xml(root, source, source, {"stylesheet": "../outside.xsl"})


if __name__ == "__main__":
    unittest.main()
