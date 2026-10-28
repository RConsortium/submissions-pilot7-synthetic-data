"""Browser views of retained workbooks and original Define-XML stylesheets."""

import html
from pathlib import Path
import shutil

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def workbook_body(path):
    """Keep sheet, row and column coordinates, merged cells and literal cell text."""
    book = load_workbook(path, data_only=False)
    sections, tables = [], []
    try:
        navigation = (
            '<div class="source-file-tabs workbook-tabs" role="tablist" aria-label="Worksheets">'
            + "".join(
                f'<button type="button" role="tab" id="sheet-tab-{n}" '
                f'aria-controls="sheet-{n}" aria-selected="{str(n == 1).lower()}" '
                f'tabindex="{0 if n == 1 else -1}">{html.escape(sheet.title)}</button>'
                for n, sheet in enumerate(book, 1)
            )
            + "</div>"
        )
        for n, sheet in enumerate(book, 1):
            populated_rows = {}
            for number, row in enumerate(sheet.iter_rows(), 1):
                values = [str(c.value) if c.value is not None else "" for c in row]
                if any(value.strip() for value in values):
                    populated_rows[number] = (row, values)
            merged, covered = {}, set()
            for region in sheet.merged_cells.ranges:
                region_rows = [
                    number for number in populated_rows
                    if region.min_row <= number <= region.max_row
                ]
                if not region_rows:
                    continue
                merged[region_rows[0], region.min_col] = (
                    len(region_rows),
                    region.max_col - region.min_col + 1,
                )
                covered.update(
                    (r, c)
                    for r in region_rows
                    for c in range(region.min_col, region.max_col + 1)
                    if (r, c) != (region_rows[0], region.min_col)
                )
            rows = []
            for number, (row, values) in populated_rows.items():
                anchor = f"sheet-{n}-row-{number}"
                sections.append(
                    {
                        "id": anchor,
                        "title": f"{sheet.title} / row {number}",
                        "text": " | ".join(values),
                    }
                )
                cells = []
                for column, cell in enumerate(row, 1):
                    if (number, column) in covered:
                        continue
                    rowspan, colspan = merged.get((number, column), (1, 1))
                    classes = ' class="workbook-bold"' if cell.font.bold else ""
                    cells.append(
                        f'<td rowspan="{rowspan}" colspan="{colspan}"{classes}>{html.escape(values[column - 1])}</td>'
                    )
                rows.append(
                    f'<tr id="{anchor}"><th scope="row" class="row-number">{number}</th>{"".join(cells)}</tr>'
                )
            headings = "".join(
                f'<th scope="col">{get_column_letter(c)}</th>'
                for c in range(1, sheet.max_column + 1)
            )
            tables.append(
                f'<section class="workbook-sheet" id="sheet-{n}" role="tabpanel" '
                f'aria-labelledby="sheet-tab-{n}" tabindex="0"{" hidden" if n != 1 else ""}>'
                f'<header class="pane-header"><h2>{html.escape(sheet.title)}</h2>'
                f'<span>{len(rows)} rows / {sheet.max_column} columns</span></header>'
                '<div class="workbook-table" role="region" '
                f'aria-label="{html.escape(sheet.title)} worksheet table" tabindex="0">'
                f'<table class="review-table"><caption class="sr-only">{html.escape(sheet.title)}</caption>'
                f'<thead><tr><th scope="col" class="row-number">Row</th>{headings}</tr></thead>'
                f'<tbody>{"".join(rows)}</tbody></table>'
                + ('<p class="empty-state">This worksheet is empty.</p>' if not rows else '')
                + '</div></section>'
            )
        controls = (
            '<div class="workbook-toolbar"><p id="workbook-status" role="status"></p>'
            + ('<button type="button" class="button" id="compare-worksheets" '
               'aria-pressed="false">Compare worksheets</button>' if len(tables) > 1 else '')
            + '</div>'
        )
        return controls + navigation + '<div class="workbook-panes">' + "".join(tables) + '</div>', sections
    finally:
        book.close()


def workbook_html(title, body):
    """Reuse dashboard table styles in a portable, sandboxed worksheet viewer."""
    docs = Path(__file__).resolve().parents[3] / "docs"
    css = (docs / "stylesheets/dashboard.css").read_text()
    script = (docs / "assets/dashboard/workbook.js").read_text()
    return (
        '<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f'<title>{html.escape(title)}</title><style>{css}</style></head>'
        '<body class="study-dashboard workbook-viewer">'
        f'<h1>{html.escape(title)}</h1>'
        '<p class="workbook-note">Blank rows are omitted; original sheet, row and column '
        'references are preserved. Formulas are shown as entered.</p>'
        f'{body}<script>{script}</script></body></html>'
    )


def render_xml(root, path, original, entry):
    """Apply the retained publisher stylesheet without restyling its HTML."""
    stylesheet = entry.get("stylesheet")
    if not stylesheet:
        return None
    source = (root / stylesheet).resolve()
    source.relative_to(root.resolve())
    retained = original.parent / source.name
    if source != retained.resolve():
        shutil.copyfile(source, retained)
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    transform = etree.XSLT(
        etree.parse(str(source), parser),
        access_control=etree.XSLTAccessControl.DENY_ALL,
    )
    view = original.with_name(original.name + ".original.html")
    view.write_bytes(bytes(transform(etree.parse(str(path), parser))))
    return view
