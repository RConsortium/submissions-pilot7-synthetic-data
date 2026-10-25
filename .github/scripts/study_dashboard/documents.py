"""Retained submission documents and exact, escaped reading-copy locations."""

import html
import json
import shutil
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=True, separators=(",", ":")) + "\n")


def local_name(element):
    return element.tag.rsplit("}", 1)[-1]


def description(element):
    text = " ".join(
        " ".join("".join(x.itertext()).split())
        for x in element
        if local_name(x) == "Description"
    ).strip()
    # Define-XML 1.0 stores the label in def:Label, not Description.
    return text or next(
        (v for k, v in element.attrib.items() if k.rsplit("}", 1)[-1] == "Label"), ""
    )


def define_metadata(path):
    tree = ET.parse(path)
    items = {x.attrib["OID"]: x for x in tree.iter() if local_name(x) == "ItemDef"}
    methods = {
        x.attrib["OID"]: description(x)
        for x in tree.iter()
        if local_name(x) == "MethodDef"
    }
    groups = {}
    for group in tree.iter():
        if local_name(group) != "ItemGroupDef":
            continue
        fields = {}
        for reference in group:
            if local_name(reference) != "ItemRef":
                continue
            item = items.get(reference.attrib.get("ItemOID"))
            if item is None:
                continue
            fields[item.attrib["Name"]] = {
                "label": description(item),
                "type": item.attrib.get("DataType"),
                "oid": item.attrib["OID"],
                "method": methods.get(reference.attrib.get("MethodOID"), ""),
            }
        groups[group.attrib.get("Name", "")] = fields
    return groups


def document_html(title, body):
    return (
        '<!doctype html><html lang="en"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>{html.escape(title)}</title><style>"
        "body{font:16px/1.7 system-ui;margin:30px auto;padding:0 24px;max-width:1000px;color:#24344a}"
        "section{border-bottom:1px solid #ddd;padding:24px 0;scroll-margin-top:20px}"
        "pre{white-space:pre-wrap;font:inherit}td,th{padding:8px;border:1px solid #ddd;vertical-align:top}"
        "table{border-collapse:collapse}a{color:#2d5f8a}:target{background:#fff5d6}"
        "</style><body>" + f"<h1>{html.escape(title)}</h1>" + body + "</body></html>"
    )


def materials(root, study, target, metadata):
    folder = target / "documents"
    folder.mkdir(exist_ok=True)
    manifest_path = root / study / "spec/source/manifest.json"
    manifest = (
        json.loads(manifest_path.read_text())
        if manifest_path.exists()
        else {"files": []}
    )
    documents = []
    entries = list(manifest.get("files", []))
    known = {e["path"] for e in entries}
    for path in sorted((root / study / "spec/define").glob("*.xml")):
        relative = path.relative_to(root).as_posix()
        if relative not in known:
            entries.append({"path": relative, "role": "published_metadata"})
    for entry in entries:
        path = root / entry["path"]
        path.resolve().relative_to(root.resolve())
        relative = path.relative_to(root)
        original = folder / relative
        if not path.is_file():
            documents.append(
                {
                    "name": path.name,
                    "path": entry["path"],
                    "available": False,
                    "role": entry.get("role"),
                    "upstream": entry.get("upstream", {}),
                    "sections": [],
                }
            )
            continue
        original.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(path, original)
        record = {
            "name": path.name,
            "path": entry["path"],
            "role": entry.get("role"),
            "upstream": entry.get("upstream", {}),
            "original": f"documents/{relative.as_posix()}",
            "available": True,
            "sections": [],
        }
        if path.suffix.lower() == ".pdf":
            for number, page in enumerate(PdfReader(path).pages, 1):
                record["sections"].append(
                    {
                        "id": f"page-{number}",
                        "title": f"Page {number}",
                        "text": page.extract_text() or "",
                    }
                )
        elif path.suffix.lower() == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=True)
            for sheet_number, sheet in enumerate(book, 1):
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(value) if value is not None else "" for value in row]
                    if any(values):
                        record["sections"].append(
                            {
                                "id": f"sheet-{sheet_number}-row-{number}",
                                "title": f"{sheet.title} / row {number}",
                                "text": " | ".join(values),
                            }
                        )
            book.close()
        elif path.name == "define-adam.xml":
            for domain, columns in metadata.items():
                for name, item in columns.items():
                    record["sections"].append(
                        {
                            "id": f"{domain}.{name}",
                            "title": f'{domain}.{name} / {item["oid"]}',
                            "text": f'Label: {item["label"]}\nData type: {item["type"]}\nMethod: {item["method"]}',
                        }
                    )
        if record["sections"]:
            body = "<p>Extracted reading copy. Verify against the linked original submission document.</p>"
            body += f'<p><a href="{html.escape(path.name)}">Open original {html.escape(path.name)}</a></p>'
            for section in record["sections"]:
                body += (
                    f'<section id="{html.escape(section["id"])}"><h2>{html.escape(section["title"])}</h2>'
                    f'<pre>{html.escape(section["text"])}</pre></section>'
                )
            reading = str(relative) + ".html"
            (folder / reading).write_text(
                document_html(path.name, body), encoding="utf-8"
            )
            record["reading"] = f"documents/{reading}"
        documents.append(record)
    save_json(target / "documents.json", documents)
    return manifest, documents
